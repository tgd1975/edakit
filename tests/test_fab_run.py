"""Stages L and F end to end (#18): `fab.run(cfg, board, out, offline=True)` on an export written for the test –
every row passes – and one fault per check built into a copy of it: the check finds it and names the part, net or
file. Each test lists every row that fails or warns, so a check that fires as well shows up.

The export is the example's input stage: J1 (JST PH, through-hole), D1 (SS34 against reverse polarity) and C2 (22 µF)
with the nets VIN, +5V and GND, on a board of 30 x 20 mm with one M3 hole and two tooling holes, every track on TOP,
no pours (`pours = []`).
The whole example has never been routed or exported, so no real export of it exists. This one has the files that
`fab.collect` leaves, with the fields the checks read: pcb.json, drc.json, bom_easyeda.json, cpl_easyeda.xlsx,
jlc_parts.json (so F-2 stays offline), and the Gerber and drill files with FlyingProbeTesting.json. pcb.json has every
designator visible (L-17). It also has the
footprint geometry of `edakit pcb dump`, the net currents of `edakit sim` – 0.5 A on all three nets, which puts
them in class POWER (0.6 mm), with an IPC-2221 width of 0.20 mm – and the board as `edakit route read` saw it (pads
with their shapes, for L-13 and L-14), written from pcb.json when the export is written. J1 is through-hole and the
example orders Economic assembly (SMD only), so J1 is not in its BOM: it is soldered by hand.

A new check of stage L or F gets its fault here: one more entry in FAULTS.
"""

import json
import math

import openpyxl
import pytest

from edakit import cli, gerber
from edakit.release import fab

MIL = 1000 / 25.4
INPUT = ("J1", "D1", "C2")
NET = {"J1.1": "VIN", "J1.2": "GND", "D1.1": "+5V", "D1.2": "VIN", "C2.1": "+5V", "C2.2": "GND"}
LCSC = {"J1": "C131337", "D1": "C8678", "C2": "C45783"}
FOOTPRINTS = {  # at 0° on TOP, mm: body (w, h), pads {number: (x, y)}, pad aperture, through-hole
    "CONN-TH_B2B-PH-K-S": ((6.0, 4.5), {"1": (-1.0, 0.0), "2": (1.0, 0.0)}, ("C", 1.4, 1.4), True),
    "SMA_L4.3-W2.6-LS5.2-RD": ((6.6, 3.0), {"1": (-2.2, 0.0), "2": (2.2, 0.0)}, ("R", 2.5, 1.7), False),
    "C0805": ((3.7, 1.9), {"1": (-0.95, 0.0), "2": (0.95, 0.0)}, ("R", 1.0, 1.3), False),
}
PLACE = {"J1": ("CONN-TH_B2B-PH-K-S", 3, 10, 90), "D1": ("SMA_L4.3-W2.6-LS5.2-RD", 13, 9, 180),
         "C2": ("C0805", 20, 9, 0)}
TRACKS = {"VIN": ["J1.1", "D1.2"], "+5V": ["D1.1", "C2.1"],
          "GND": ["J1.2", (3, 13.5), (22.5, 13.5), (22.5, 9), "C2.2"]}
WIDTH, DRILL = 0.6, 1.0            # tracks, J1's plated holes (mm)
HOLE = (26, 16, 3.2)               # M3, a circle on the outline layer
TOOLING = [(2.5, 17.5), (27.5, 2.5)]
TOOLING_D, TOOLING_OPEN = 1.152, 1.152 + 2 * 0.148
PLUS = (10.8, 10.9)                # "+" on the top silkscreen, next to D1's anode
CATHODE = (15.2, 10.9)             # 1.9 mm above D1.1, D1's cathode
BOARD = f"""# The input stage of the example on a board of its own (tests/test_fab_run.py)
view = "top"
pours = []
edge_ok = ["J1"]

[outline]
rect = [30, 20]

[[holes]]
x = {HOLE[0]}
y = {HOLE[1]}
d = {HOLE[2]}
keep = 3.2

[tooling]
holes = {[list(t) for t in TOOLING]}

[place]
""" + "".join(f'{r} = {{ x = {x}, y = {y}, rot = {rot}, side = "top" }}\n' for r, (_, x, y, rot) in PLACE.items()) + """
[near]
C2 = "D1"

[[loops]]
a = "J1.1"
b = "D1.2"
max_mm = 10
what = "input"

[routing]
hand_solder_keep_mm = 0.5        # L-18: J1 is soldered by hand (bom: false), its joints are on BOTTOM
"""


def turn(p, rot):
    a = math.radians(rot)
    c, s = round(math.cos(a)), round(math.sin(a))
    return p[0] * c - p[1] * s, p[0] * s + p[1] * c


def cut(project):
    """The example's graph cut down to its input stage, and the board file of that stage."""
    g = project.graph()
    gg = g["graph"]
    gg["components"] = [c for c in gg["components"] if c["designator"] in INPUT]
    gg["nets"] = [{**n, "pins": [p for p in n["pins"] if p.split(".")[0] in INPUT]} for n in gg["nets"]]
    gg["nets"] = [n for n in gg["nets"] if n["pins"]]
    gg["layout"] = {"edge": ["J1"]}
    project.save_graph(g)
    (project.root / "hardware" / "boards" / "main.toml").write_text(BOARD)


class Layer(list):
    """Gerber items in mm: ("flash", (kind, a, b), x, y) and ("draw", width, points, tag)."""

    def text(self):
        aps, body = {}, []
        um = lambda v: round(v * 1e6)  # noqa: E731

        def code(ap):
            return aps.setdefault(ap, 10 + len(aps))
        for item in self:
            if item[0] == "flash":
                body += [f"D{code(item[1])}*", f"X{um(item[2])}Y{um(item[3])}D03*"]
            else:
                _, w, pts, _ = item
                body += [f"D{code(('C', w, w))}*", f"X{um(pts[0][0])}Y{um(pts[0][1])}D02*"]
                body += [f"X{um(x)}Y{um(y)}D01*" for x, y in pts[1:]]
        head = ["%FSLAX46Y46*%", "%MOMM*%"] + [f"%ADD{c}{k},{a:.6f}" + ("" if k == "C" else f"X{b:.6f}") + "*%"
                                              for (k, a, b), c in aps.items()]
        return "\n".join(head + body + ["M02*"]) + "\n"

    def drop(self, tag):
        self[:] = [i for i in self if not (i[0] == "draw" and i[3] == tag)]


def drill(holes):
    """Excellon in mm: holes (x, y, d), slots (x0, y0, x1, y1, d)."""
    tools = {}
    for h in holes:
        tools.setdefault(h[-1], len(tools) + 1)
    out = ["M48", "METRIC"] + [f"T{t:02d}C{d:.3f}" for d, t in tools.items()] + ["%"]
    for d, t in tools.items():
        out.append(f"T{t:02d}")
        out += [f"X{h[0]:.3f}Y{h[1]:.3f}" + (f"G85X{h[2]:.3f}Y{h[3]:.3f}" if len(h) == 5 else "")
                for h in holes if h[-1] == d]
    return "\n".join(out + ["M30"]) + "\n"


class Export:
    """The files of one gate run, in memory: a fault changes them, write() puts them where fab.run reads them."""

    def __init__(self):
        self.at, comps, self.geo, probe = {}, [], {}, []
        cu = {k: Layer() for k in gerber.FILES if k not in ("pth", "vias", "npth", "probe")}
        self.pth, self.npth = [], [(x, y, TOOLING_D) for x, y in TOOLING]
        self.vias = []                 # (x, y, d): EasyEDA writes them into the PTH file and into a file of their own
        for ref, (fp, x, y, rot) in PLACE.items():
            (w, h), pads, (k, a, b), tht = FOOTPRINTS[fp]
            self.geo[ref] = {"footprint": fp, "bbox": [-w / 2 * MIL, -h / 2 * MIL, w / 2 * MIL, h / 2 * MIL],
                             "pads": {n: {"x": px * MIL, "y": py * MIL, "net": NET[f"{ref}.{n}"], "tht": tht}
                                      for n, (px, py) in pads.items()}}
            a, b = (b, a) if rot % 180 == 90 else (a, b)
            rows = []
            for n, p in pads.items():
                dx, dy = turn(p, rot)
                px, py = self.at[f"{ref}.{n}"] = (x + dx, y + dy)
                rows.append({"PadNumber": n, "X": px * MIL, "Y": py * MIL, "Net": NET[f"{ref}.{n}"],
                             "Layer": fab.MULTI if tht else 1, "tht": tht})
                for layer in ("top", "bottom") if tht else ("top",):
                    cu[layer].append(("flash", (k, a, b), px, py))
                    cu[f"{layer}_mask"].append(("flash", (k, a + 0.1, b + 0.1), px, py))
                if tht:
                    self.pth.append((px, py, DRILL))
                probe.append([f"{ref}_{n}", px * MIL, py * MIL, NET[f"{ref}.{n}"], "THT" if tht else "SMD", "T",
                              a * MIL, b * MIL])
            comps.append({"designator": ref, "footprintName": fp, "x": x * MIL, "y": y * MIL, "rotation": rot,
                          "layer": 1, "pads": rows})
        lines = []
        for net, path in TRACKS.items():
            pts = [self.at[p] if isinstance(p, str) else p for p in path]
            cu["top"].append(("draw", WIDTH, pts, net))
            lines += [{"Net": net, "Layer": 1, "LineWidth": WIDTH * MIL, "StartX": p[0] * MIL, "StartY": p[1] * MIL,
                       "EndX": q[0] * MIL, "EndY": q[1] * MIL} for p, q in zip(pts, pts[1:])]
        for x, y in TOOLING:
            for mask in ("top_mask", "bottom_mask"):
                cu[mask].append(("flash", ("C", TOOLING_OPEN, TOOLING_OPEN), x, y))
        cu["top_silk"] += [("draw", 0.15, [(PLUS[0] - 0.4, PLUS[1]), (PLUS[0] + 0.4, PLUS[1])], "+"),
                           ("draw", 0.15, [(PLUS[0], PLUS[1] - 0.4), (PLUS[0], PLUS[1] + 0.4)], "+")]
        cu["outline"].append(("draw", 0.1, [(0, 0), (30, 0), (30, 20), (0, 20), (0, 0)], "outline"))
        ring = [(HOLE[0] + HOLE[2] / 2 * math.cos(math.pi * k / 12), HOLE[1] + HOLE[2] / 2 * math.sin(math.pi * k / 12))
                for k in range(25)]
        cu["outline"].append(("draw", 0.1, ring, "hole"))
        self.gerber = cu
        self.pcb = {"components": comps, "lines": lines,
                    "strings": [{"Text": "+", "X": PLUS[0] * MIL, "Y": PLUS[1] * MIL, "Layer": 3}],
                    "designators": {ref: True for ref in PLACE}}
        self.probe = {"pins": {"fields": ["PIN_NAME", "PIN_X", "PIN_Y", "NET_NAME", "PIN_TYPE", "LAYER", "PAD_SIZEX",
                                          "PAD_SIZEY"], "rows": probe}}
        self.drc = []
        self.netlist_diff = None       # the dry runs of `pcb parts` and `pcb nets` (only with an Import Changes item)
        self.bom = [{"reference": r, "lcsc": LCSC[r], "value": v, "quantity": 1}
                    for r, v in (("J1", ""), ("D1", "SS34"), ("C2", "22uF"))]
        self.jlc = {"C8678": {"componentModelEn": "SS34", "componentLibraryType": "base", "stockCount": 52000},
                    "C45783": {"componentModelEn": "CL21A226MAQNNNE", "componentLibraryType": "base",
                               "stockCount": 310000}}
        self.cpl = [{"Designator": r, "Footprint": fp, "Mid X": x, "Mid Y": y, "Layer": "T", "Rotation": rot,
                     "SMD": "No" if FOOTPRINTS[fp][3] else "Yes"} for r, (fp, x, y, rot) in PLACE.items()]
        self.currents = {"nets": {n: {"i_max_a": 0.5, "scenario": "written for the test"} for n in TRACKS}}
        # pad shape per footprint at 0° as `route read` gives it (mil); None: `route read` never ran. read_off:
        # {ref: dx mm} – the part stood there when the board was read
        kinds = {"C": "ELLIPSE", "R": "RECT"}
        self.shapes = {fp: [kinds[k], a * MIL, b * MIL] for fp, (_, _, (k, a, b), _) in FOOTPRINTS.items()}
        self.read_off = {}
        self.lost = set()              # keys of gerber.FILES the export leaves out
        self.silk_memo = None          # silk_strings_main.json of `silk apply`; None: it never ran
        self.kept = None               # {silk key: Layer}: the footprint print `silk apply` saved; None: none saved

    def board_read(self):
        """<board>_board.json of `edakit route read`, from the PCB as pcb.json has it."""
        comps, pads = [], []
        for c in self.pcb["components"]:
            ref, dx = c["designator"], self.read_off.get(c["designator"], 0) * MIL
            comps.append({"ref": ref, "x": c["x"] + dx, "y": c["y"], "rot": c["rotation"], "layer": c["layer"]})
            pads += [{"id": f"{ref}e{p['PadNumber']}", "ref": ref, "num": p["PadNumber"], "net": p["Net"],
                      "x": p["X"] + dx, "y": p["Y"], "rot": c["rotation"], "layer": p["Layer"],
                      "shape": self.shapes[PLACE.get(ref, [c["footprintName"]])[0]],
                      "hole": [DRILL * MIL] if p["tht"] else None,
                      "hole_off": [0, 0]} for p in c["pads"]]
        edge = [(0, 0), (30 * MIL, 0), (30 * MIL, 20 * MIL), (0, 20 * MIL)]
        return {"pcb": "main", "components": comps, "pads": pads, "outlines": [edge]}

    def at_mil(self, ref):
        return tuple(v * MIL for v in self.at[ref])

    def comp(self, ref):
        return next(c for c in self.pcb["components"] if c["designator"] == ref)

    def pad(self, ref):
        des, _, num = ref.partition(".")
        return next(p for p in self.comp(des)["pads"] if p["PadNumber"] == num)

    def move(self, ref, dx):
        c = self.comp(ref)
        c["x"] += dx * MIL
        for p in c["pads"]:
            p["X"] += dx * MIL

    def turn(self, ref, rot):
        c = self.comp(ref)
        c["rotation"] = (c["rotation"] + rot) % 360
        for p in c["pads"]:
            dx, dy = turn((p["X"] - c["x"], p["Y"] - c["y"]), rot)
            p["X"], p["Y"] = c["x"] + dx, c["y"] + dy

    def line(self, net, layer, pts, width=WIDTH):
        self.gerber["top" if layer == 1 else "bottom"].append(("draw", width, pts, net))
        self.pcb["lines"] += [{"Net": net, "Layer": layer, "LineWidth": width * MIL, "StartX": a[0] * MIL,
                               "StartY": a[1] * MIL, "EndX": b[0] * MIL, "EndY": b[1] * MIL} for a, b in zip(pts, pts[1:])]

    def width(self, net, mm):
        for ln in self.pcb["lines"]:
            if ln["Net"] == net:
                ln["LineWidth"] = mm * MIL

    def write(self, root):
        art = root / "build" / "edakit"
        (art / "sim").mkdir(parents=True, exist_ok=True)
        (art / "route").mkdir()
        for f, data in (("pcb_geometry_main.json", self.geo), ("sim/net_currents.json", self.currents),
                        ("route/main_board.json", self.shapes and self.board_read())):
            if data is not None:       # None: `pcb dump`, `edakit sim` or `route read` never ran
                (art / f).write_text(json.dumps(data))
        if self.silk_memo is not None:
            (art / "silk_strings_main.json").write_text(json.dumps(self.silk_memo))
        if self.kept is not None:
            (art / "silk_footprints_main").mkdir()
            for key, layer in self.kept.items():
                (art / "silk_footprints_main" / gerber.FILES[key]).write_text(layer.text())
        d = root / "build" / "fab" / "main"
        (d / "gerber").mkdir(parents=True)
        for name, data in (("pcb", self.pcb), ("drc", self.drc), ("bom_easyeda", self.bom), ("jlc_parts", self.jlc),
                           ("netlist_diff", self.netlist_diff)):
            if data is not None:
                (d / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False))
        wb = openpyxl.Workbook()
        ws = wb.active
        head = ["Designator", "Footprint", "Mid X", "Mid Y", "Layer", "Rotation", "SMD"]
        ws.append(head)
        for r in self.cpl:
            ws.append([r[k] for k in head])
        wb.save(d / "cpl_easyeda.xlsx")
        g = d / "gerber"
        files = {key: layer.text() for key, layer in self.gerber.items()}
        files.update(pth=drill(self.pth + self.vias), npth=drill(self.npth), probe=json.dumps(self.probe))
        if self.vias:
            files["vias"] = drill(self.vias)
        for key, text in files.items():
            if key not in self.lost:
                (g / gerber.FILES[key]).write_text(text)
        return d.parent


def run(project, fault=None):
    cut(project)
    ex = Export()
    if fault:
        fault(ex, project)
    out = ex.write(project.root)
    cfg = project.config()
    return fab.run(cfg, cfg.board(), out, offline=True)


def findings(rows):
    return sorted((r["id"], r["result"]) for r in rows if r["result"] not in ("pass", "info"))


CHECKS = ["L-1", "L-2", "L-3", "L-4", "L-5", "L-7", "L-8", "L-9", "L-10", "L-11", "L-12", "L-13", "L-14", "L-15", "L-16",
          "L-17", "L-18", "F-1", "F-2", "F-3", "F-5"]
ON_REQUEST = ["F-6", "L-19"]       # run only when the board file asks for them ([order], [looks])


def test_the_clean_export_passes_every_check(project):
    rows = run(project)
    assert findings(rows) == []
    assert sorted({r["id"] for r in rows if r["stage"] in "LF"}, key=CHECKS.index) == CHECKS   # no [order]: no F-6
    assert not (project.root / "build" / "fab" / "main" / "order.md").exists()
    by = {(r["id"], r["check"]): r["actual"] for r in rows}
    assert by["L-11", "short current loops: path over tracks per pad pair within the limit"] == "J1.1-D1.2 7.8 mm"
    assert by["L-9", "hole takes the lead at nominal size (diagonal or edge, or manufacturer's value)"] == "2 holes"
    assert by["L-8", "silkscreen '+'/'−' = role of the nearest polarised pad"] == \
        "1 signs correct: 1 as PCB text, 0 drawn in the silkscreen"     # the text's strokes are that text
    paste_row = next(r for r in rows if r["id"] == "L-14")
    assert (paste_row["actual"], paste_row["expected"]) == (
        "4 SMD pads", ">= 0.400 mm on TOP ([release] assembly paste_gap_mm)")     # J1 is through-hole


def test_edakit_fab_offline_on_the_clean_export(project, monkeypatch, capsys):
    cut(project)
    out = Export().write(project.root)
    monkeypatch.chdir(project.root)
    monkeypatch.delenv("EDAKIT_CONFIG", raising=False)
    assert cli.main(["fab", "--offline", "--out", str(out)]) == 0
    rows = json.loads(capsys.readouterr().out.splitlines()[-1])
    assert findings(rows) == []


def test_edakit_order_writes_the_bom_of_f1(project, monkeypatch):
    """#41: the assembler's BOM of `edakit order` is byte for byte the one F-1 writes – C2 and D1; J1 is soldered by
    hand and goes to a cart."""
    run(project)
    monkeypatch.chdir(project.root)
    monkeypatch.delenv("EDAKIT_CONFIG", raising=False)
    assert cli.main(["order", "--out", str(project.root / "order")]) == 0
    bom = (project.root / "order" / "bom_main.csv").read_bytes()
    assert bom == (project.root / "build" / "fab" / "main" / "bom_main.csv").read_bytes()
    assert [line.split(",")[-1] for line in bom.decode().splitlines()] == ["LCSC Part #", "C45783", "C8678"]
    assert "C131337" in (project.root / "order" / "cart_lcsc.csv").read_text()


def board_file(project, old, new):
    f = project.root / "hardware" / "boards" / "main.toml"
    assert old in f.read_text()
    f.write_text(f.read_text().replace(old, new))


def graph(project, change):
    g = project.graph()
    change(g)
    project.save_graph(g)


# the DRC item EasyEDA writes while the import view has not been opened since schematic parts changed (#67)
IMPORT = {"errorType": "Netlist Error", "ruleName": "Import Changes", "obj1": {"typeName": "Schematic Netlist"},
          "obj2": {"typeName": "PCB Netlist"}}
CLEARANCE = {"errorType": "Clearance", "ruleName": "Safe Spacing"}
NO_DIFF = {"parts": {}, "nets": []}


def import_pending(diff, *more):
    def fault(ex, project):
        ex.drc += [IMPORT, *more]
        ex.netlist_diff = diff
    return fault


def short(ex, project):
    ex.gerber["top"].append(("draw", 0.3, [ex.at["J1.1"], ex.at["J1.2"]], "short"))


def detour(ex, project):
    a, b = ex.at["J1.1"], ex.at["D1.2"]
    pts = [a, (a[0], 5), (b[0], 5), b]
    ex.pcb["lines"] = [ln for ln in ex.pcb["lines"] if ln["Net"] != "VIN"] + [
        {"Net": "VIN", "Layer": 1, "LineWidth": WIDTH * MIL, "StartX": p[0] * MIL, "StartY": p[1] * MIL,
         "EndX": q[0] * MIL, "EndY": q[1] * MIL} for p, q in zip(pts, pts[1:])]


def slot(ex, project):
    x, y, _ = ex.pth[0]
    ex.pth[0] = (x, y - 0.375, x, y + 0.375, 0.45)


def crowd(ex, project):
    """C2 fixed left of D1, its GND pad 0.09 mm from D1's VIN pad (anode) – PCB and plan agree, and the boxes reach
    0.16 mm into each other, within the plan's silk_overlap_mm."""
    board_file(project, "C2 = { x = 20,", "C2 = { x = 8.01,")
    board_file(project, 'edge_ok = ["J1"]', 'edge_ok = ["J1"]\nsilk_overlap_mm = 0.2')
    ex.move("C2", 8.01 - 20)


def paste(ex, project):
    """C2 fixed left of D1, its GND pad 0.30 mm from D1's VIN pad (anode): above the clearance, below the paste gap of
    0.4 mm on the assembled top side. The boxes stay 0.05 mm apart."""
    board_file(project, "C2 = { x = 20,", "C2 = { x = 7.8,")
    ex.move("C2", 7.8 - 20)


def bare_bottom(ex, project):
    """GND poured on BOTTOM by the board file, but the BOTTOM Gerber holds only J1's two pads: a pour never filled."""
    board_file(project, "pours = []\n", "")
    f = project.root / "hardware" / "boards" / "main.toml"
    f.write_text(f.read_text() + '\n[[pours]]\nnet = "GND"\nlayers = ["BOTTOM"]\n')


def loose_copper(ex, project):
    """A copper stroke of 3 x 3 mm with round ends (16 mm²) on TOP that touches no pad, via or hole."""
    ex.gerber["top"].append(("draw", 3.0, [(9, 4.5), (12, 4.5)], "loose"))


def j1_pads(mm):
    """J1's round pads (Ø 1.4 mm on 1.0 mm holes, a ring of 0.2 mm) at another diameter on both copper layers."""
    def fault(ex, project):
        for layer in ("top", "bottom"):
            ex.gerber[layer][:] = [("flash", ("C", mm, mm), *i[2:]) if i[:2] == ("flash", ("C", 1.4, 1.4)) else i
                                   for i in ex.gerber[layer]]
    return fault


def via(ex, x, y, pad=0.6096, hole=0.3048):
    """A via as edakit routes it ([24, 12] mil), tented: drill files and a pad on both copper layers."""
    ex.vias.append((x, y, hole))
    for layer in ("top", "bottom"):
        ex.gerber[layer].append(("flash", ("C", pad, pad), x, y))


def plus_at(x, y, text=True):
    """The "+" of the export moved to (x, y): drawn in the top silkscreen as two crossing strokes, the way a footprint
    draws it, and as a PCB text unless text is False (a footprint's own print)."""
    def fault(ex, project):
        ex.gerber["top_silk"].drop("+")
        ex.gerber["top_silk"] += [("draw", 0.15, [(x - 0.4, y), (x + 0.4, y)], "+"),
                                  ("draw", 0.15, [(x, y - 0.4), (x, y + 0.4)], "+")]
        ex.pcb["strings"] = [{**s, "X": x * MIL, "Y": y * MIL} for s in ex.pcb["strings"]] if text else []
    return fault


def silk_line(*pts):
    """A 0.15 mm line of print on the top silkscreen."""
    def fault(ex, project):
        ex.gerber["top_silk"].append(("draw", 0.15, list(pts), "line"))
    return fault


ACROSS_C2 = ((18, 9), (20, 9))     # across C2.1's mask opening (1.1 x 1.4 mm): 0.165 mm² of print on it


def own_print(ex, project):
    """`silk apply` ran – its memo, and the print it found before placing anything, saved – and then a line of its
    own print lies across C2.1: a part moved after it, say."""
    ex.silk_memo = {"strings": ["s1"], "fills": []}
    ex.kept = {k: Layer(ex.gerber[k]) for k in ("top_silk", "bottom_silk")}
    silk_line(*ACROSS_C2)(ex, project)


def via_at_joint(ex, project):
    """A GND via 0.40 mm (copper to copper) left of J1.1 (VIN), with its piece of GND track on TOP – J1 sits on TOP, so
    only the via is on J1's solder side."""
    via(ex, 1.6, 9)
    ex.line("GND", 1, [(1.6, 9), (1.6, 6)])


def track_at_joint(ex, project):
    """A GND track on BOTTOM from J1.2 that passes 0.20 mm below J1.1 (VIN), J1's solder side – tracks do not count
    for L-18. It starts on J1.2, so L-15 sees no loose copper."""
    y = 9 - 0.7 - 0.2 - WIDTH / 2
    ex.line("GND", 2, [(3, 11), (1.5, 11), (1.5, y), (5, y)])


def assembled_j1(ex, project):
    graph(project, lambda g: project.comp(g, "J1").pop("bom"))
    ex.jlc["C131337"] = {"componentModelEn": "B2B-PH-K-S(LF)(SN)", "componentLibraryType": "extend",
                         "stockCount": 41000}


REMARK = "No V-cut. Use our 2 NPTH tooling holes (1.152 mm). Tabs only on the right edge X30 Y0-20, none elsewhere. Ask us first."
PLACEHOLDER = (18.0, 3.0)          # the order-number placeholder on the top silkscreen, bottom right


def ordered(ex, project, order=None, placeholder=True, h_mm=1.0):
    """#40: [order] in the board file (mark = "specify" and REMARK unless order says otherwise), and the placeholder
    JLCJLCJLCJLC on the top silkscreen – in pcb.json and as [[silk.texts]] of h_mm."""
    order = {"mark": "specify", "remark": REMARK, **(order or {})}
    text = "\n[order]\n" + "".join(f"{k} = {json.dumps(v, ensure_ascii=False)}\n" for k, v in order.items()
                                   if k != "tab_zones")
    if "tab_zones" in order:
        text += "tab_zones = [" + ", ".join(f"{{ from = {list(a)}, to = {list(b)} }}" for a, b in order["tab_zones"]) + "]\n"
    if placeholder:
        text += (f'\n[[silk.texts]]\ntext = "JLCJLCJLCJLC"\nx = {PLACEHOLDER[0]}\ny = {PLACEHOLDER[1]}\nh_mm = {h_mm}\n'
                 'side = "top"\n')
        ex.pcb["strings"].append({"Text": "JLCJLCJLCJLC", "X": PLACEHOLDER[0] * MIL, "Y": PLACEHOLDER[1] * MIL, "Layer": 3})
    f = project.root / "hardware" / "boards" / "main.toml"
    f.write_text(f.read_text() + text)


def edge_track(ex, project):
    """A tab zone along the left edge, and a track of 0.2 mm whose copper runs 0.45 mm from that edge."""
    ordered(ex, project, {"tab_zones": [((0, 0), (0, 20))]})
    ex.gerber["top"].append(("draw", 0.2, [(0.55, 3), (0.55, 15)], "edge"))


def no_tooling(ex, project):
    ordered(ex, project)
    board_file(project, "[tooling]\nholes", "# [tooling]\n# holes")


TEST = (15.0, 3.0)                 # a PCB text "TEST" on the top silkscreen (#81)


def visible(project, side="top", silk="", plus=True):
    """#81: [looks] visible = side, the [silk] table `silk`, and – plus – the export's "+" at D1 as a [[silk.texts]]
    entry of the top side, so that the board file puts it there."""
    text = f'\n[looks]\nvisible = "{side}"\n\n[silk]\n{silk}\n'
    if plus:
        text += f'\n[[silk.texts]]\ntext = "+"\nx = {PLUS[0]}\ny = {PLUS[1]}\nside = "top"\n'
    f = project.root / "hardware" / "boards" / "main.toml"
    f.write_text(f.read_text() + text)


def visible_designator(ex, project):
    """The designators belong on BOTTOM, TOP is the visible side, and C2 shows its designator there – a part placed
    again after silk apply hid it."""
    visible(project, silk='designators = "bottom"\nhide = ["top"]')
    ex.pcb["designators"].update(J1=False, D1=False)


def stray_text(ex, project):
    """A PCB text "TEST" on the visible TOP side that no line of the board file puts there."""
    visible(project)
    ex.pcb["strings"].append({"Text": "TEST", "X": TEST[0] * MIL, "Y": TEST[1] * MIL, "Layer": 3})


def thickness(mm, order=None):
    """[order] in the board file (as ordered) on a board of thickness_mm = mm."""
    def fault(ex, project):
        ordered(ex, project, order)
        board_file(project, 'edge_ok = ["J1"]', f'edge_ok = ["J1"]\nthickness_mm = {mm}')
    return fault


FAULTS = [
    # L-1
    ("a pad on another net", lambda ex, p: ex.pad("D1.1").update(Net="VIN"),
     [("L-1", "fail", "D1.1: VIN instead of +5V")]),
    ("a part missing", lambda ex, p: ex.pcb["components"].remove(ex.comp("C2")),
     [("L-1", "fail", "missing ['C2']")]),
    ("a part too many", lambda ex, p: ex.pcb["components"].append({**ex.comp("C2"), "designator": "C9"}),
     [("L-1", "fail", "extra ['C9']")]),
    ("a graph pin without a pad", lambda ex, p: ex.comp("C2")["pads"].remove(ex.pad("C2.2")),
     [("L-1", "fail", "C2.2: pad missing in the footprint (GND)")]),
    # L-2
    ("a short across the connector", short, [("L-2", "fail", "GND + VIN (C2_2, D1_2, J1_1, J1_2)")]),
    ("an open net", lambda ex, p: ex.gerber["top"].drop("+5V"), [("L-2", "fail", "+5V: 2 groups")]),
    ("a connected pad missing from the files", lambda ex, p: ex.probe["pins"]["rows"].pop(4),
     [("L-2", "fail", "C2.1")]),
    ("a copper file missing", lambda ex, p: ex.lost.add("top"),
     [("L-2", "fail", "Gerber_TopLayer.GTL missing"), ("F-5", "fail", "Gerber_TopLayer.GTL"),
      ("L-15", "fail", "Gerber_TopLayer.GTL missing"), ("L-16", "fail", "Gerber_TopLayer.GTL missing")]),
    ("the probe file missing", lambda ex, p: ex.lost.add("probe"),
     [("L-2", "fail", "FlyingProbeTesting.json missing"), ("F-5", "fail", "FlyingProbeTesting.json"),
      ("L-15", "fail", "FlyingProbeTesting.json missing")]),
    # L-3
    ("a DRC finding", lambda ex, p: ex.drc.append(CLEARANCE), [("L-3", "fail", "Clearance (Safe Spacing): 1")]),
    ("the import acknowledgement pending", import_pending(NO_DIFF),
     [("L-3", "fail", "EasyEDA waits for the import acknowledgement: open Design → Import Changes once")]),
    ("a property of D1 differs between schematic and PCB", import_pending({"parts": {"props": ["D1"]}, "nets": []}),
     [("L-3", "fail", "properties differ from the schematic: D1 → edakit pcb parts --write")]),
    ("a pad net differs between graph and PCB", import_pending({"parts": {}, "nets": [["D1.2", "+5V", "VIN"]]}),
     [("L-3", "fail", "pad nets: D1.2 +5V instead of VIN → edakit pcb nets --write")]),
    ("an Import Changes item, schematic and PCB not compared", import_pending(None),
     [("L-3", "fail", "schematic and PCB not compared")]),
    ("an Import Changes item and a clearance error", import_pending(NO_DIFF, CLEARANCE),
     [("L-3", "fail", "Clearance (Safe Spacing): 1"), ("L-3", "fail", "EasyEDA waits for the import acknowledgement")]),
    # L-4
    ("another footprint", lambda ex, p: ex.comp("C2").update(footprintName="C0603"),
     [("L-4", "fail", "C2: C0603 instead of C0805")]),
    # L-5
    ("a power net at 3 mil", lambda ex, p: ex.width("VIN", 3 / MIL),
     [("L-5", "fail", "VIN: 0.08 < 0.20 mm"), ("L-5", "warn", "VIN: 0.08 < 0.60 mm")]),
    ("a power net below its class width", lambda ex, p: ex.width("VIN", 0.4),
     [("L-5", "warn", "VIN: 0.40 < 0.60 mm")]),
    ("no simulated currents", lambda ex, p: setattr(ex, "currents", None),
     [("L-5", "fail", "net_currents.json missing (edakit sim)")]),
    # L-7
    ("a fixed part moved by 1 mm", lambda ex, p: ex.move("C2", 1),
     [("L-7", "fail", "C2: (21.00, 9.00, 0°, top) instead of (20.00, 9.00, 0°, top)")]),
    ("a fixed part turned by 90 degrees", lambda ex, p: ex.turn("C2", 90),
     [("L-7", "fail", "C2: (20.00, 9.00, 90°, top) instead of (20.00, 9.00, 0°, top)"),
      ("F-3", "fail", "C2: rotation 0 differs from the PCB")]),
    ("a plan finding", lambda ex, p: board_file(p, 'C2 = "D1"', 'C2 = "J1"'),
     [("L-7", "fail", "C2: 12.9 mm from J1 (near_mm 6)")]),
    ("no footprint geometry", lambda ex, p: setattr(ex, "geo", None),
     [("L-7", "fail", "pcb_geometry_main.json missing (edakit pcb dump)")]),
    # L-8
    ("a minus sign at the anode", lambda ex, p: ex.pcb["strings"][0].update(Text="−"),
     [("L-8", "fail", "'−' at (10.8, 10.9) mm: nearest pad D1.2 is anode")]),
    ("a plus sign on a side without polarised pads",
     lambda ex, p: ex.pcb["strings"].append({**ex.pcb["strings"][0], "Layer": 4}),
     [("L-8", "fail", "'+' at (10.8, 10.9) mm: no polarised pad on this side")]),
    ("a plus sign at the cathode", plus_at(*CATHODE),           # a cathode had no sign: compared with the anode
     [("L-8", "fail", "'+' at (15.2, 10.9) mm: nearest pad D1.1 is cathode")]),
    ("a footprint's own plus sign at the cathode", plus_at(*CATHODE, text=False),   # only the Gerber has it
     [("L-8", "fail", "'+' at (15.2, 10.9) mm: nearest pad D1.1 is cathode")]),
    # L-9
    ("a through-hole part without lead data", lambda ex, p: p.save_part("C131337", {
        k: v for k, v in p.part("C131337").items() if k != "lead"}), [("L-9", "fail", "J1")]),
    ("holes too small for the lead", lambda ex, p: ex.pth.__setitem__(slice(None), [h[:2] + (0.5,) for h in ex.pth]),
     [("L-9", "fail", "J1.1: hole 0.50 mm for diagonal 0.71 mm")]),
    ("a plated slot narrower than 0.5 mm", slot,        # 1.2 mm long in J1's round 1.4 mm pad: 0.1 mm at its ends
     [("L-9", "fail", "J1.1: slot 0.45 x 1.20 mm"), ("L-9", "warn", "pad 1: 0.45 x 1.20 mm"),
      ("L-16", "fail", "J1.1: 0.100 mm on TOP and BOTTOM (hole 0.45 x 1.20 mm)")]),
    ("a through-hole pad without its hole", lambda ex, p: ex.pth.pop(),     # the bottom pad of J1.2 loses GND too
     [("L-9", "fail", "J1.2: no hole"), ("L-2", "fail", "GND: 2 groups")]),
    ("the drill file missing", lambda ex, p: ex.lost.add("pth"),
     [("L-2", "fail", "Drill_PTH_Through.DRL missing"), ("L-9", "fail", "Drill_PTH_Through.DRL missing"),
      ("F-5", "fail", "Drill_PTH_Through.DRL"), ("L-16", "fail", "Drill_PTH_Through.DRL missing")]),
    # L-10
    ("a pad away from its library position", lambda ex, p: ex.pad("C2.1").update(X=ex.at_mil("C2.1")[0] + 0.2 * MIL),
     [("L-10", "fail", "C2.1: 0.20 mm away from its footprint position")]),
    # L-11
    ("a current loop over its limit", detour, [("L-11", "fail", "J1.1-D1.2 (input): 15.8 mm > 10.0 mm")]),
    ("a loop's part missing in the PCB", lambda ex, p: ex.pcb["components"].remove(ex.comp("J1")),
     [("L-1", "fail", "missing ['J1']"), ("L-11", "fail", "J1.1-D1.2 (input): J1.1 not in the PCB")]),
    ("a loop's pads on different nets in the PCB", lambda ex, p: ex.pad("D1.2").update(Net="+5V"),
     [("L-1", "fail", "D1.2: +5V instead of VIN"), ("L-11", "fail", "J1.1-D1.2 (input): different nets (VIN, +5V)")]),
    # L-12
    ("the circuit's edge part not fixed", lambda ex, p: board_file(p, "J1 = { x = 3", "# J1 = { x = 3"),
     [("L-12", "fail", "J1: not in hardware/boards/main.toml [place] – the circuit wants it at the board edge")]),
    # L-13
    ("two parts with their pads closer than the clearance", crowd,
     [("L-13", "fail", "C2.2 (GND) – D1.2 (VIN): 0.090 mm apart on TOP, below the clearance 0.165 mm"),
      ("L-14", "fail", "C2.2 (GND) – D1.2 (VIN): 0.090 mm apart on TOP, below the paste gap 0.400 mm")]),
    ("the board never read by route read", lambda ex, p: setattr(ex, "shapes", None),
     [("L-13", "fail", "build/edakit/route/main_board.json missing (edakit route read)"),
      ("L-14", "fail", "build/edakit/route/main_board.json missing (edakit route read)"),
      ("L-18", "fail", "build/edakit/route/main_board.json missing (edakit route read)")]),
    ("a part moved since the board was read", lambda ex, p: ex.read_off.update(C2=1),
     [("L-13", "fail", "C2.1, C2.2 moved or changed since: edakit route read"),
      ("L-14", "fail", "C2.1, C2.2 moved or changed since: edakit route read"),
      ("L-18", "fail", "C2.1, C2.2 moved or changed since: edakit route read")]),
    # L-14
    ("two SMD parts with their pads closer than the paste gap", paste,
     [("L-14", "fail", "C2.2 (GND) – D1.2 (VIN): 0.300 mm apart on TOP, below the paste gap 0.400 mm")]),
    # L-15
    ("a poured layer left bare", bare_bottom,
     [("L-15", "warn", "BOTTOM: 1 % copper, largest copper-free circle Ø 19.0 mm at (13.2, 10.1) mm")]),
    ("copper that touches no pad, via or hole", loose_copper,
     [("L-15", "warn", "TOP: 1 piece without pad, via or hole: 16.2 mm² around (10.5, 4.5) mm")]),
    # L-16
    ("a pad ring below JLCPCB's minimum", j1_pads(1.25),
     [("L-16", "fail", "J1.1: 0.125 mm on TOP and BOTTOM (hole Ø 1.00 mm); J1.2: 0.125 mm on TOP and BOTTOM")]),
    ("a pad ring below JLCPCB's capability value", j1_pads(1.38),
     [("L-16", "warn", "pad 1, 2: 0.190 mm on TOP and BOTTOM (hole Ø 1.00 mm)")]),
    ("a via ring below 0.05 mm", lambda ex, p: via(ex, 8, 13.5, pad=0.38),     # on TOP the GND track is its ring
     [("L-16", "fail", "via at (8.00, 13.50): 0.037 mm on BOTTOM (hole Ø 0.30 mm)")]),
    ("two vias closer than 0.2 mm", lambda ex, p: (via(ex, 8, 13.5), via(ex, 8.45, 13.5)),
     [("L-16", "fail", "via at (8.00, 13.50) – via at (8.45, 13.50): 0.145 mm apart, below 0.20 mm")]),
    ("a via closer than 0.45 mm to a component hole", lambda ex, p: via(ex, 3, 12.05),     # J1 is soldered by hand
     [("L-16", "fail", "J1.2 – via at (3.00, 12.05): 0.398 mm apart, below 0.45 mm"),
      ("L-18", "fail", "J1.2 (GND) – via (no track, net unknown) at (3.00, 12.05): 0.05 mm apart on BOTTOM")]),
    # L-17
    ("edakit's own print on a mask opening", own_print,
     [("L-17", "fail", "TOP: 0.165 mm² in 1 spot: (19.1, 9.0) mm at C2.1")]),
    ("edakit printed, but its print cannot be told from the footprints'",
     lambda ex, p: setattr(ex, "silk_memo", {"strings": ["s1"], "fills": []}),
     [("L-17", "fail", "silk_strings_main.json exists, silk_footprints_main/ does not")]),
    ("a footprint's print on a mask opening", silk_line(*ACROSS_C2),
     [("L-17", "warn", "TOP: 0.165 mm² in 1 spot: (19.1, 9.0) mm at C2.1")]),
    ("print beyond the board outline", silk_line((29, 15), (31, 15)),
     [("L-17", "warn", "TOP: 0.157 mm² beyond it in 1 spot: around (30.5, 15.0) mm")]),
    ("a part without a visible designator", lambda ex, p: ex.pcb["designators"].update(C2=False),
     [("L-17", "warn", "C2 on the top without a visible designator")]),
    ("the designators not read with the PCB", lambda ex, p: ex.pcb.pop("designators"),
     [("L-17", "fail", "pcb.json without the designators")]),
    ("a text height below the fab's minimum",
     lambda ex, p: board_file(p, 'edge_ok = ["J1"]', 'edge_ok = ["J1"]\n\n[silk]\nlabel_h_mm = 0.8'),
     [("L-17", "warn", "label_h_mm = 0.8 mm")]),
    # L-18
    ("a via of another net next to a hand-soldered joint", via_at_joint,
     [("L-18", "fail", "J1.1 (VIN) – via GND at (1.60, 9.00): 0.40 mm apart on BOTTOM, the solder side of J1, below "
       "hand_solder_keep_mm 0.5")]),
    ("a via without a track next to a hand-soldered joint", lambda ex, p: via(ex, 1.6, 9),
     [("L-18", "fail", "J1.1 (VIN) – via (no track, net unknown) at (1.60, 9.00): 0.40 mm")]),
    # L-19 (#81) – only with [looks] in the board file
    ("a designator shown on the visible side", visible_designator,
     [("L-19", "fail", "TOP (visible): C2 shows its designator – [silk] designators is bottom")]),
    ("a text on the visible side the board file does not put there", stray_text,
     [("L-19", "fail", "TOP (visible): text 'TEST' at (15.0, 3.0) mm is neither the board label, a designator, a "
       "group label nor a [[silk.texts]] entry of this side")]),
    ("the designators not read with the PCB, the visible side without them",
     lambda ex, p: (visible(p, silk='designators = "bottom"'), ex.pcb.pop("designators")),
     [("L-17", "fail", "pcb.json without the designators"), ("L-19", "fail", "pcb.json without the designators")]),
    # F-1
    ("another LCSC number in EasyEDA's BOM", lambda ex, p: ex.bom[2].update(lcsc="C15850"),
     [("F-1", "fail", "C2: graph C45783, EasyEDA C15850")]),
    ("an LCSC number not in the parts database", lambda ex, p: (p.root / "hardware" / "parts" / "C45783.json").unlink(),
     [("F-1", "fail", "C45783")]),
    # F-2
    ("a part not in JLCPCB's catalogue", lambda ex, p: ex.jlc.pop("C8678"),
     [("F-2", "fail", "D1 C8678: not in the catalogue (pre-order only)")]),
    ("a part out of stock", lambda ex, p: ex.jlc["C45783"].update(stockCount=3),
     [("F-2", "fail", "C2 C45783: 3 in stock, 5 needed")]),
    ("the catalogue not reachable", lambda ex, p: ex.jlc.update(C8678={"error": "timed out"}),
     [("F-2", "fail", "D1 C8678: catalogue not reachable (timed out)")]),
    ("the catalogue names another part (#38)", lambda ex, p: ex.jlc["C8678"].update(componentModelEn="SS14"),
     [("F-2", "fail", "D1 C8678: JLCPCB names it SS14, the parts entry SS34")]),
    ("the catalogue lists another class (#38)", lambda ex, p: ex.jlc["C8678"].update(componentLibraryType="expand"),
     [("F-2", "warn", "D1 C8678: JLCPCB lists it as Extended Part, the parts entry as Basic Part")]),
    # F-3
    ("a part missing from the pick-and-place file", lambda ex, p: ex.cpl.pop(1), [("F-3", "fail", "D1: 0x")]),
    ("a part twice in the pick-and-place file", lambda ex, p: ex.cpl.append(dict(ex.cpl[2])),
     [("F-3", "fail", "C2: 2x")]),
    ("a part on the other side in the pick-and-place file", lambda ex, p: ex.cpl[2].update(Layer="B"),
     [("F-3", "fail", "C2: side B differs from the PCB"), ("F-3", "fail", "Bottom SMD 1")]),
    ("a through-hole part in the BOM of an SMD-only assembly", assembled_j1, [("F-3", "fail", "Top THT 1")]),
    # F-5
    ("a Gerber file missing", lambda ex, p: ex.gerber.pop("outline"),
     [("F-5", "fail", "Gerber_BoardOutlineLayer.GKO"), ("L-17", "fail", "Gerber_BoardOutlineLayer.GKO missing")]),
    ("a silkscreen file missing", lambda ex, p: ex.lost.add("bottom_silk"),
     [("F-5", "fail", "Gerber_BottomSilkscreenLayer.GBO"), ("L-8", "fail", "Gerber_BottomSilkscreenLayer.GBO missing"),
      ("L-17", "fail", "Gerber_BottomSilkscreenLayer.GBO missing")]),
    ("no outline in the outline file", lambda ex, p: ex.gerber["outline"].clear(),
     [("F-5", "fail", "none"), ("L-17", "fail", "no closed outline in Gerber_BoardOutlineLayer.GKO")]),
    ("the outline off the plan", lambda ex, p: ex.gerber["outline"].__setitem__(0, (
        "draw", 0.1, [(0, 0), (31, 0), (31, 20), (0, 20), (0, 0)], "outline")),
     [("F-5", "fail", "31.00 x 20.00 mm at (0.00, 0.00)")]),
    ("a mounting hole missing", lambda ex, p: ex.gerber["outline"].drop("hole"), [("F-5", "fail", "missing 1")]),
    ("a tooling hole missing from the NPTH file", lambda ex, p: ex.npth.pop(),
     [("F-5", "fail", "1 NPTH holes, missing 1")]),
    ("the mask closed over a tooling hole", lambda ex, p: ex.gerber["top_mask"].pop(),
     [("F-5", "fail", "top_mask: opening at (27.5, 2.5) smaller than Ø 1.448 mm")]),
    ("copper in a tooling hole's opening",     # the dot is no pad either: what the drill leaves of it is loose
     lambda ex, p: ex.gerber["top"].append(("flash", ("C", 0.3, 0.3), 3.15, 17.5)),
     [("F-5", "fail", "top: copper in the opening at (2.5, 17.5)"), ("L-15", "warn", "TOP: 1 piece without pad")]),
    # F-6 (#40) – only with [order] in the board file
    ("a remark longer than 200 characters", lambda ex, p: ordered(ex, p, {"remark": "x" * 201}),
     [("F-6", "fail", "201 characters")]),
    ("a remark that is not ASCII", lambda ex, p: ordered(ex, p, {"remark": "Tabs only where copper is ≥ 1 mm away"}),
     [("F-6", "fail", "not ASCII: ≥")]),
    ("mark specify without the placeholder", lambda ex, p: ordered(ex, p, placeholder=False),
     [("F-6", "fail", "no JLCJLCJLCJLC on a silkscreen layer")]),
    ("the placeholder too small", lambda ex, p: ordered(ex, p, h_mm=0.6),   # below L-17's 1.0 mm for edakit's texts too
     [("F-6", "fail", "JLCJLCJLCJLC 0.60 mm high"), ("L-17", "warn", "'JLCJLCJLCJLC' = 0.6 mm")]),
    ("the placeholder kept with mark default", lambda ex, p: ordered(ex, p, {"mark": "default"}),
     [("F-6", "warn", 'JLCJLCJLCJLC at (18.0, 3.0) on the top silkscreen, but mark = "default"')]),
    ("Economic without tooling holes", no_tooling,
     [("F-6", "warn", "no tooling holes ([tooling]) – JLCPCB adds its own")]),
    ("copper in a tab zone", edge_track,       # the track touches no pad: loose copper for L-15 too
     [("F-6", "fail", "copper 0.45 mm from the edge at (0.0, "), ("L-15", "warn", "TOP: 1 piece without pad")]),
    ("a board thickness JLCPCB does not make", thickness(1.57),         # 62 mil: JLCPCB's form has 1.6 mm
     [("F-6", "fail", "1.57 mm (thickness_mm) – not a thickness of JLCPCB's form")]),
    ("white print on a white mask", lambda ex, p: ordered(ex, p, {"mask": "white", "silkscreen": "white"}),
     [("F-6", "fail", 'silkscreen = "white" on a white mask – JLCPCB prints black on it')]),
]


@pytest.mark.parametrize("fault, expected", [pytest.param(f, e, id=name) for name, f, e in FAULTS])
def test_each_fault_is_found_by_its_check(project, fault, expected):
    rows = run(project, fault)
    assert findings(rows) == sorted((i, r) for i, r, _ in expected)
    for cid, result, text in expected:
        assert any(r["id"] == cid and r["result"] == result and text in r["actual"] for r in rows), \
            [(r["result"], r["actual"]) for r in rows if r["id"] == cid]


def test_a_waived_pad_pair_is_listed_with_its_reason(project):
    def waived(ex, p):
        crowd(ex, p)
        graph(p, lambda g: p.comp(g, "C2").update(waive={"L-13": "test: close on purpose"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-13", "waived"), ("L-14", "fail")]          # a waiver holds for its own check
    w = next(r for r in rows if r["result"] == "waived")
    assert w["reason"] == "test: close on purpose" and "C2.2 (GND) – D1.2 (VIN): 0.090 mm" in w["check"]


def test_a_waived_paste_gap_is_listed_with_its_reason(project):
    def waived(ex, p):
        paste(ex, p)
        graph(p, lambda g: next(n for n in g["graph"]["nets"] if n["name"] == "VIN").update(
            waive={"L-14": "test: paste reduced on D1"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-14", "waived")]
    w = next(r for r in rows if r["result"] == "waived")
    assert w["reason"] == "test: paste reduced on D1" and "C2.2 (GND) – D1.2 (VIN): 0.300 mm" in w["check"]


def test_no_paste_gap_on_a_side_that_is_not_assembled(project):
    def bottom_only(ex, p):
        paste(ex, p)
        f = p.root / "edakit.toml"
        f.write_text(f.read_text().replace('sides = ["top"]', 'sides = ["bottom"]'))   # F-3 fails, not L-14
    rows = run(project, bottom_only)
    assert not [r for r in rows if r["id"] == "L-14" and r["result"] != "pass"]
    assert next(r for r in rows if r["id"] == "L-14")["actual"] == "0 SMD pads"


def test_copper_per_side_of_the_clean_export(project):
    """L-15 on the clean export: nothing poured (pours = []), so share and bare circle are information; the 3 pieces of
    TOP (the three nets' tracks with their pads) and J1's two pads on BOTTOM all touch a pad."""
    rows = {r["check"]: (r["result"], r["actual"]) for r in run(project) if r["id"] == "L-15"}
    assert rows == {
        "TOP: copper share and largest copper-free circle (no pour)":
            ("info", "TOP: 6 % copper, largest copper-free circle Ø 8.2 mm at (5.0, 4.5) mm"),
        "TOP: every copper piece touches a pad, via or plated hole": ("pass", "TOP: 3 pieces"),
        "BOTTOM: copper share and largest copper-free circle (no pour)":
            ("info", "BOTTOM: 1 % copper, largest copper-free circle Ø 19.0 mm at (13.2, 10.1) mm"),
        "BOTTOM: every copper piece touches a pad, via or plated hole": ("pass", "BOTTOM: 2 pieces")}


def test_copper_per_side_settings(project):
    """[copper] of the board file: max_bare_mm lets the bare BOTTOM pass, min_cover warns below its share."""
    def settings(ex, p):
        bare_bottom(ex, p)
        f = p.root / "hardware" / "boards" / "main.toml"
        f.write_text(f.read_text() + "\n[copper]\nmax_bare_mm = 20\nmin_cover = { TOP = 10 }\n")
    rows = run(project, settings)
    assert findings(rows) == [("L-15", "warn")]
    by = {r["check"]: (r["result"], r["actual"]) for r in rows if r["id"] == "L-15"}
    assert by["BOTTOM: no copper-free circle above 20 mm inside the pour outline (poured: GND)"][0] == "pass"
    assert by["TOP: copper share at least 10 % ([copper] min_cover)"] == ("warn", "TOP: 6 % copper")


def test_rings_and_hole_spacing_of_the_clean_export(project):
    """L-16 on the clean export: J1's two holes of 1.0 mm in pads of 1.4 mm, 2 mm apart – a ring of 0.2 mm is kept."""
    rows = [(r["result"], r["actual"]) for r in run(project) if r["id"] == "L-16"]
    assert rows == [("pass", "2 holes (narrowest ring 0.200 mm, J1.1), 0 vias"), ("pass", "closest 1.000 mm (J1.1 – J1.2)")]


def test_a_waived_annular_ring_is_listed_with_its_reason(project):
    def waived(ex, p):
        j1_pads(1.25)(ex, p)
        graph(p, lambda g: p.comp(g, "J1").update(waive={"L-16": "test: JLCPCB confirmed the ring"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-16", "waived"), ("L-16", "waived")]                  # one row per hole
    w = [r for r in rows if r["result"] == "waived"]
    assert {r["reason"] for r in w} == {"test: JLCPCB confirmed the ring"}
    assert [(r["check"], r["actual"]) for r in w] == [
        ("annular ring: J1.1: 0.125 mm on TOP and BOTTOM (hole Ø 1.00 mm)", "0.125 mm"),
        ("annular ring: J1.2: 0.125 mm on TOP and BOTTOM (hole Ø 1.00 mm)", "0.125 mm")]


def test_a_minus_sign_at_the_cathode_is_right(project):
    def minus(ex, p):
        ex.pcb["strings"] = [{"Text": "−", "X": CATHODE[0] * MIL, "Y": CATHODE[1] * MIL, "Layer": 3}]
    rows = run(project, minus)                  # the drawn "+" stays at the anode
    assert findings(rows) == []
    assert next(r for r in rows if r["id"] == "L-8")["actual"] == \
        "2 signs correct: 1 as PCB text, 1 drawn in the silkscreen"


def test_a_cross_in_the_artwork_is_not_read(project):
    def art(ex, p):
        plus_at(*CATHODE, text=False)(ex, p)
        (p.root / "star.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
                                         '<rect x="0" y="0" width="10" height="10" fill="black"/></svg>')
        board_file(p, 'edge_ok = ["J1"]', 'edge_ok = ["J1"]\n\n[[silk.art]]\nsvg = "star.svg"\n'
                   f"x = {CATHODE[0]}\ny = {CATHODE[1]}\nwidth_mm = 3")
    rows = run(project, art)
    assert findings(rows) == []
    assert not [r for r in rows if r["id"] == "L-8"]                    # no sign left to compare


def test_a_waived_sign_is_listed_with_its_reason(project):
    def waived(ex, p):
        plus_at(*CATHODE, text=False)(ex, p)
        graph(p, lambda g: p.comp(g, "D1").update(waive={"L-8": "test: the library's print, the roles are right"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-8", "waived")]
    w = next(r for r in rows if r["result"] == "waived")
    assert w["reason"] == "test: the library's print, the roles are right"
    assert "'+' at (15.2, 10.9) mm: nearest pad D1.1 is cathode" in w["check"]


def test_plus_signs_in_a_silkscreen_file(tmp_path):
    lay = Layer([("draw", 0.15, [(0, 1), (1.2, 1)], "+"), ("draw", 0.15, [(0.6, 0.4), (0.6, 1.6)], "+"),   # a "+"
                 ("draw", 0.15, [(5, 0), (6, 0), (6, 1), (5, 1), (5, 0)], "frame"),       # corners meet at their ends
                 ("draw", 0.15, [(10, 0), (10.5, 0.866)], "X"), ("draw", 0.15, [(10.5, 0), (10, 0.866)], "X"),  # 60°
                 ("draw", 0.15, [(20, 0), (20.2, 0)], "short"), ("draw", 0.15, [(20.1, -0.1), (20.1, 0.1)], "short"),
                 ("draw", 0.15, [(30, 0), (33, 0)], "long"), ("draw", 0.15, [(31.5, -1.5), (31.5, 1.5)], "long")])
    (tmp_path / "silk.gbr").write_text(lay.text())
    assert [tuple(round(v, 3) for v in p) for p in gerber.plus_signs(tmp_path / "silk.gbr")] == [(0.6, 1.0)]


def l17(rows):
    return {r["check"]: (r["result"], r["actual"]) for r in rows if r["id"] == "L-17"}


def test_silkscreen_of_the_clean_export(project):
    assert l17(run(project)) == {
        "edakit's own print (designators, texts, artwork of silk apply) off the solder-mask openings":
            ("pass", "edakit printed nothing on this board (no silk apply)"),
        "TOP: other print (footprints, texts by hand) off the solder-mask openings": ("pass", "TOP: 0.000 mm²"),
        "BOTTOM: other print (footprints, texts by hand) off the solder-mask openings": ("pass", "BOTTOM: 0.000 mm²"),
        "TOP: silkscreen inside the board outline": ("pass", "TOP: 0.000 mm² beyond it"),
        "BOTTOM: silkscreen inside the board outline": ("pass", "BOTTOM: 0.000 mm² beyond it"),
        "designators of the top side read with the PCB": ("pass", "3 parts: 3 with a visible designator"),
        "text heights edakit prints at least 1.0 mm (JLCPCB minimum text height)": ("pass", "label_h_mm 1 mm"),
    }


def test_silk_apply_run_again_leaves_only_the_footprints_print(project):
    """After `silk apply` the footprint print it saved covers the print on the opening: a warning for C2, no fail."""
    def again(ex, p):
        silk_line(*ACROSS_C2)(ex, p)
        ex.silk_memo = {"strings": [], "fills": []}
        ex.kept = {k: Layer(ex.gerber[k]) for k in ("top_silk", "bottom_silk")}
    rows = run(project, again)
    assert findings(rows) == [("L-17", "warn")]
    assert l17(rows)["TOP: edakit's own print (designators, texts, artwork of silk apply) off the solder-mask "
                     "openings"] == ("pass", "TOP: 0.000 mm²")
    assert l17(rows)["C2: silkscreen on its solder-mask openings – JLCPCB clips it there"][0] == "warn"


def test_a_group_label_or_unlabelled_covers_a_part(project):
    def grouped(ex, p):
        ex.pcb["designators"].update(J1=False, C2=False, D1=False)
        ex.pcb["strings"].append({"Text": "IN", "X": 10 * MIL, "Y": 15 * MIL, "Layer": 3})
        board_file(p, 'edge_ok = ["J1"]', 'edge_ok = ["J1"]\n\n[silk]\ngroups = { IN = ["J1", "D1"] }\n'
                   'unlabelled = { C2 = "no room next to D1" }')
    rows = run(project, grouped)
    assert findings(rows) == []
    assert l17(rows)["designators of the top side read with the PCB"] == (
        "pass", "3 parts: 0 with a visible designator, 2 under a group label, unlabelled on purpose: "
                "C2 (no room next to D1)")


def test_the_acknowledgement_alone_leaves_the_other_kinds_passing(project):
    rows = run(project, import_pending(NO_DIFF))
    other, sync = [r for r in rows if r["id"] == "L-3"]
    assert (other["check"], other["result"]) == ("EasyEDA DRC without findings (other kinds)", "pass")
    assert (sync["check"], sync["result"]) == ("schematic and PCB in sync (Import Changes)", "fail")
    assert "it shows no changes" in sync["actual"]


def test_every_check_has_a_fault():
    assert {cid for _, _, e in FAULTS for cid, _, _ in e} == set(CHECKS + ON_REQUEST)


def test_the_order_form_of_the_clean_export(project):
    """#40: with [order] F-6 passes, and order.md holds every field of the order form – the class and side from
    [release] assembly, the tooling holes, the mark at the placeholder, the remark, bake, the free outline stretches."""
    rows = run(project, ordered)
    assert findings(rows) == []
    f6 = [r for r in rows if r["id"] == "F-6"]
    assert {r["result"] for r in f6} == {"pass", "info"}
    assert len(REMARK) <= 200 and any(r["actual"] == f"{len(REMARK)} characters, ASCII" for r in f6)
    assert {(r["check"], r["actual"]) for r in f6} >= {
        ("board thickness JLCPCB makes (FR-4, 2 layers)", "1.6 mm (thickness_mm)"),       # JLCPCB's defaults
        ("silkscreen colour JLCPCB prints on the solder mask", "white on a green mask")}
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "| PCB Qty | 5 |\n| PCB Thickness | 1.6 mm |\n| PCB Color | Green |\n| Silkscreen | White |\n" \
           "| Surface Finish | LeadFree HASL |" in sheet                            # in the order of JLCPCB's form
    for text in ("| Layers | 2 |", "| Dimensions | 30.00 x 20.00 mm (Gerber outline) |", "| PCB Qty | 5 |",
                 "| Surface Finish | LeadFree HASL |",
                 "| Mark on PCB | Order Number (Specify Position) – JLCJLCJLCJLC at (18.0, 3.0) on the top silkscreen |",
                 REMARK, "| PCBA Type | Economic ([release] assembly.sides top: one side) |",
                 "| Assembly Side | Top Side |", "| PCBA Qty | 5 |", "| Tooling holes | Added by Customer (2 x Ø 1.152 mm) |",
                 "| Bake Components | check the BOM tool – no stock snapshot (edakit stock) |",
                 "| BOM / CPL | bom_main.csv (2 lines) / cpl_main.csv |", "| Gerber | gerber_main.zip |",
                 "(30.0, 0.0) → (30.0, 20.0), 20.0 mm", "| J1 | C131337 B2B-PH-K-S(LF)(SN) | lcsc | C131337 |"):
        assert text in sheet, text


def test_the_order_form_lists_only_the_free_edges(project):
    """#40: copper 0.45 mm from the left edge – the tab zone there fails, the right edge is still free."""
    rows = run(project, edge_track)
    row = next(r for r in rows if r["id"] == "F-6" and r["result"] == "fail")
    assert row["actual"].startswith("copper 0.45 mm from the edge at (0.0, ") and row["expected"] == ">= 1.00 mm (tab_copper_mm)"
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "(30.0, 0.0) → (30.0, 20.0), 20.0 mm" in sheet and "| (0.0, 0.0) → (0.0, 20.0) | 0.45 mm" in sheet


def test_the_order_form_bakes_an_msl5_part(project):
    """#40: bake "Yes" when the stock snapshot gives a BOM part MSL 5 or more."""
    def snapshot(ex, p):
        ordered(ex, p)
        rows = {"C8678": {"date": "2026-10-27", "listed": True, "msl": 5}, "C45783": {"date": "2026-10-27", "listed": True,
                                                                               "msl": 1}}
        (p.root / "hardware" / "stock.json").write_text(json.dumps({"date": "2026-10-27", "parts": rows}))
    run(project, snapshot)
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "| Bake Components | Yes – D1 C8678 MSL 5 (stock snapshot 2026-10-27) |" in sheet


def test_the_order_form_names_thickness_and_colours(project):
    """Step 1.7: order.md takes the board thickness from thickness_mm and the colours from [order] mask and
    silkscreen; F-6 passes a thickness of JLCPCB's form and the print JLCPCB puts on the mask."""
    rows = run(project, thickness(1.0, {"mask": "black"}))
    assert findings(rows) == []
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "| PCB Thickness | 1 mm |\n| PCB Color | Black |\n| Silkscreen | White |\n" in sheet
    f6 = {r["check"]: (r["result"], r["actual"]) for r in rows if r["id"] == "F-6"}
    assert f6["board thickness JLCPCB makes (FR-4, 2 layers)"] == ("pass", "1 mm (thickness_mm)")
    assert f6["silkscreen colour JLCPCB prints on the solder mask"] == ("pass", "white on a black mask")


def test_the_order_form_of_a_white_mask(project):
    """Step 1.7: on a white mask JLCPCB prints black – the silkscreen without a key of its own."""
    rows = run(project, lambda ex, p: ordered(ex, p, {"mask": "white"}))
    assert findings(rows) == []
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "| PCB Color | White |\n| Silkscreen | Black |\n" in sheet
    assert ("pass", "black on a white mask") in {(r["result"], r["actual"]) for r in rows if r["id"] == "F-6"}


def test_the_order_form_of_a_bare_board(project):
    """A board nobody assembles ([release] assembly.sides = [], every part bom: false) gets its order form too, PCB
    only – the assembly side of the sheet is not asked for."""
    def bare(ex, p):
        ordered(ex, p)
        graph(p, lambda g: [c.update(bom=False) for c in g["graph"]["components"]])
        f = p.root / "edakit.toml"
        assert 'sides = ["top"]' in f.read_text()
        f.write_text(f.read_text().replace('sides = ["top"]', "sides = []"))
    rows = run(project, bare)
    assert findings(rows) == []
    sheet = (project.root / "build" / "fab" / "main" / "order.md").read_text()
    assert "## PCB Assembly\n\nNone – no part of this board is assembled (PCB only)." in sheet


# --- a declared visible side (#61): stage L adds the side's numbers as a row of its own, and warns above a limit

def looks(project, keys, change=None):
    def fault(ex, p):
        f = p.root / "hardware" / "boards" / "main.toml"
        f.write_text(f.read_text() + f"\n[looks]\n{keys}\n")
        if change:
            change(ex, p)
    return [r for r in run(project, fault) if r["id"] == "L"]


def test_stage_l_reports_the_visible_side(project):
    row, = looks(project, 'visible = "bottom"')                                   # every track of the export is on TOP
    assert (row["result"], row["actual"], row["expected"]) == (
        "info", "BOTTOM (visible): 0.00 m of track, 0 vias (TOP 0.04 m, BOTTOM 0.00 m)", "information")


def test_stage_l_warns_above_the_limits_of_the_visible_side(project):
    def vias(ex, p):
        ex.vias += [(10, 15, 0.3), (12, 15, 0.3)]
    rows = looks(project, 'visible = "top"\nmax_track_m = 0.03\nmax_vias = 1', vias)       # 39.7 mm of track on TOP
    assert [(r["result"], r["actual"], r["expected"]) for r in rows] == [
        ("info", "TOP (visible): 0.04 m of track, 2 vias (TOP 0.04 m, BOTTOM 0.00 m)", "<= 0.03 m, <= 1 vias"),
        ("warn", "0.04 m of track", "<= 0.03 m"), ("warn", "2 vias", "<= 1 vias")]


def test_stage_l_within_the_limits_of_the_visible_side(project):
    rows = looks(project, 'visible = "top"\nmax_track_m = 0.1\nmax_vias = 2', lambda ex, p: ex.vias.append((10, 15, 0.3)))
    assert [(r["result"], r["actual"]) for r in rows] == [
        ("info", "TOP (visible): 0.04 m of track, 1 via (TOP 0.04 m, BOTTOM 0.00 m)")]


def test_a_waived_hand_soldered_joint_is_listed_with_its_reason(project):
    def waived(ex, p):
        via_at_joint(ex, p)
        graph(p, lambda g: next(n for n in g["graph"]["nets"] if n["name"] == "VIN").update(
            waive={"L-18": "test: the via is tented"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-18", "waived")]
    w = next(r for r in rows if r["result"] == "waived")
    assert w["reason"] == "test: the via is tented" and "J1.1 (VIN) – via GND at (1.60, 9.00)" in w["check"]


def test_hand_solder_keep_off_gives_no_row(project):
    def off(ex, p):
        via_at_joint(ex, p)
        board_file(p, "hand_solder_keep_mm = 0.5", "hand_solder_keep_mm = 0")
    rows = run(project, off)
    assert findings(rows) == [] and not [r for r in rows if r["id"] == "L-18"]


def test_the_clean_export_names_the_hand_soldered_joints(project):
    row = next(r for r in run(project) if r["id"] == "L-18")
    assert (row["result"], row["actual"], row["expected"]) == (
        "pass", "2 joints of 1 parts, 0 vias", ">= 0.5 mm ([routing] hand_solder_keep_mm)")


def test_a_track_at_a_hand_soldered_joint_does_not_count(project):
    rows = run(project, track_at_joint)
    assert findings(rows) == [] and next(r for r in rows if r["id"] == "L-18")["result"] == "pass"


# --- the visible side carries only the texts the board file puts there (#81, L-19)

def l19(rows):
    return [(r["result"], r["actual"]) for r in rows if r["id"] == "L-19"]


def test_the_visible_side_of_the_proposal(project):
    """#81's finding on the export: designators on BOTTOM, TOP visible, a PCB text "TEST" added – the three
    designators, "TEST" and the "+" (until it is a [[silk.texts]] entry) are findings; the [looks] row stays."""
    def fault(ex, p):
        visible(p, silk='designators = "bottom"\nhide = ["top"]', plus=False)
        ex.pcb["strings"].append({"Text": "TEST", "X": TEST[0] * MIL, "Y": TEST[1] * MIL, "Layer": 3})
    rows = run(project, fault)
    assert findings(rows) == [("L-19", "fail")] * 5
    assert sorted(a for _, a in l19(rows)) == [
        "TOP (visible): C2 shows its designator – [silk] designators is bottom",
        "TOP (visible): D1 shows its designator – [silk] designators is bottom",
        "TOP (visible): J1 shows its designator – [silk] designators is bottom",
        "TOP (visible): text '+' at (10.8, 10.9) mm is neither the board label nor a [[silk.texts]] entry of this side",
        "TOP (visible): text 'TEST' at (15.0, 3.0) mm is neither the board label nor a [[silk.texts]] entry of this "
        "side"]
    assert [r["actual"] for r in rows if r["id"] == "L"] == [
        "TOP (visible): 0.04 m of track, 0 vias (TOP 0.04 m, BOTTOM 0.00 m)"]


def test_a_visible_side_as_the_board_file_puts_it(project):
    """TOP visible and the designators side: the designators, a group label, the board label of that side and the
    [[silk.texts]] are the texts the board file puts there."""
    def fault(ex, p):
        board_file(p, 'view = "top"', 'view = "top"\nrevision = "1"')
        visible(p, silk='groups = { IN = ["J1", "D1"] }')
        ex.pcb["strings"] += [{"Text": t, "X": x * MIL, "Y": y * MIL, "Layer": 3}
                              for t, x, y in (("IN", 10, 15), ("C2", 20, 12), ("LED bar rev 1", 5, 2))]
    rows = run(project, fault)
    assert findings(rows) == []
    assert l19(rows) == [("pass", "TOP: 4 PCB texts, each one the board file puts there")]


@pytest.mark.parametrize("text, first", [
    ("LED bar rev 1", ("pass", "BOTTOM: 1 PCB text, each one the board file puts there")),
    ("LED bar rev 0", ("fail", "BOTTOM (visible): text 'LED bar rev 0' at (5.0, 2.0) mm is neither the board label nor "
                               "a [[silk.texts]] entry of this side"))])
def test_the_board_label_counts_on_its_own_side(project, text, first):
    """BOTTOM visible, the board label printed there (board_label_side): it passes; an old revision left on the board
    is no text the board file puts there. The designators of TOP are not on the visible side."""
    def fault(ex, p):
        board_file(p, 'view = "top"', 'view = "top"\nrevision = "1"')
        visible(p, "bottom", silk='board_label_side = "bottom"')
        ex.pcb["strings"].append({"Text": text, "X": 5 * MIL, "Y": 2 * MIL, "Layer": 4})
    assert l19(run(project, fault)) == [first, ("pass", "0 parts on the bottom, none shows its designator")]


def test_a_waived_designator_on_the_visible_side(project):
    def waived(ex, p):
        visible_designator(ex, p)
        graph(p, lambda g: p.comp(g, "C2").update(waive={"L-19": "test: the label is wanted there"}))
    rows = run(project, waived)
    assert findings(rows) == [("L-19", "waived")]
    w = next(r for r in rows if r["result"] == "waived")
    assert w["reason"] == "test: the label is wanted there"
    assert "C2 shows its designator – [silk] designators is bottom" in w["check"]
